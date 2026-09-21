"""
Management command para importar y sincronizar el Plan de Cuentas de la Empresa Olivícola
desde el archivo Excel 'Plan_de_Cuentas_Empresa_Olivicola.xlsx'.

Uso:
    python manage.py importar_plan_cuentas
    python manage.py importar_plan_cuentas --archivo ruta/al/archivo.xlsx --limpiar
"""
import os
import openpyxl
from django.core.management.base import BaseCommand
from django.conf import settings
from django.db import transaction

from apps.finanzas.models import CuentaContable, Cuenta
from apps.inventario.models import CategoriaInsumo, Insumo
from apps.core.models import CentroDeCosto


class Command(BaseCommand):
    help = 'Importa y estructura el Plan de Cuentas para Empresa Olivícola desde Excel'

    def add_arguments(self, parser):
        parser.add_argument(
            '--archivo',
            type=str,
            default='Plan_de_Cuentas_Empresa_Olivicola.xlsx',
            help='Ruta al archivo Excel con el Plan de Cuentas'
        )
        parser.add_argument(
            '--limpiar',
            action='store_true',
            help='Elimina todas las cuentas contables previas antes de importar'
        )
        parser.add_argument(
            '--sin-mapeo',
            action='store_true',
            help='Omite el enlace automático con Cajas, Insumos y Centros de Costo'
        )

    def handle(self, *args, **options):
        archivo = options['archivo']
        if not os.path.isabs(archivo):
            archivo = os.path.join(settings.BASE_DIR, archivo)

        if not os.path.exists(archivo):
            self.stderr.write(self.style.ERROR(f"[ERROR] No se encontro el archivo: {archivo}"))
            return

        self.stdout.write(self.style.NOTICE(f"[INFO] Leyendo archivo: {archivo}"))

        wb = openpyxl.load_workbook(archivo, data_only=True)
        if 'Plan de Cuentas' not in wb.sheetnames:
            self.stderr.write(self.style.ERROR("[ERROR] La hoja 'Plan de Cuentas' no existe en el archivo Excel."))
            return

        sheet = wb['Plan de Cuentas']

        raw_rows = []
        for r in range(2, sheet.max_row + 1):
            cod_val = sheet.cell(row=r, column=1).value
            nom_val = sheet.cell(row=r, column=2).value
            sal_val = sheet.cell(row=r, column=3).value
            cla_val = sheet.cell(row=r, column=4).value
            not_val = sheet.cell(row=r, column=5).value

            if not cod_val or not str(cod_val).strip():
                continue

            raw_rows.append({
                'codigo': str(cod_val).strip(),
                'nombre': str(nom_val or '').strip(),
                'saldo': str(sal_val or '').strip(),
                'clase': str(cla_val or '').strip(),
                'nota': str(not_val or '').strip() if not_val else ''
            })

        self.stdout.write(f"[INFO] Filas detectadas en el Excel: {len(raw_rows)}")

        all_codes = {item['codigo'] for item in raw_rows}

        def find_parent_code(code):
            p = code.split('.')
            if len(p) != 5:
                return None
            candidates = []
            # 1. Subpart en p[4], ej: 1.1.5.02.010101 -> 1.1.5.02.010000
            if len(p[4]) == 6 and not p[4].startswith('0000') and p[4] != '000000':
                candidates.append(f"{p[0]}.{p[1]}.{p[2]}.{p[3]}.{p[4][:2]}0000")
            # 2. Nivel 4: 1.1.1.01.000000
            if p[4] != '000000':
                candidates.append(f"{p[0]}.{p[1]}.{p[2]}.{p[3]}.000000")
            # 3. Nivel 3: 1.1.1.00.000000
            if p[3] != '00' or p[4] != '000000':
                candidates.append(f"{p[0]}.{p[1]}.{p[2]}.00.000000")
            # 4. Nivel 2: 1.1.0.00.000000
            if p[2] != '0' or p[3] != '00' or p[4] != '000000':
                candidates.append(f"{p[0]}.{p[1]}.0.00.000000")
            # 5. Nivel 1: 1.0.0.00.000000
            if p[1] != '0' or p[2] != '0' or p[3] != '00' or p[4] != '000000':
                candidates.append(f"{p[0]}.0.0.00.000000")

            for cand in candidates:
                if cand in all_codes and cand != code:
                    return cand
            return None

        def determine_level(code):
            p = code.split('.')
            if len(p) != 5:
                return 1
            if p[1] == '0' and p[2] == '0' and p[3] == '00' and p[4] == '000000':
                return 1
            if p[2] == '0' and p[3] == '00' and p[4] == '000000':
                return 2
            if p[3] == '00' and p[4] == '000000':
                return 3
            if p[4] == '000000':
                return 4
            if len(p[4]) == 6 and not p[4].startswith('0000') and p[4].endswith('0000'):
                return 5
            if len(p[4]) == 6 and not p[4].startswith('0000') and not p[4].endswith('0000'):
                return 6
            return 5

        def normalize_class(clase_str, code):
            c_low = clase_str.lower()
            if 'regularizadora' in c_low:
                return CuentaContable.ClaseCuenta.ACTIVO_REGULARIZADORA
            if 'activo' in c_low:
                return CuentaContable.ClaseCuenta.ACTIVO
            if 'pasivo' in c_low:
                return CuentaContable.ClaseCuenta.PASIVO
            if 'patrimonio' in c_low or 'resultado del ejercicio' in c_low or 'resultados acumulados' in c_low:
                return CuentaContable.ClaseCuenta.PATRIMONIO_NETO
            if 'positivo' in c_low or 'ingreso' in c_low:
                return CuentaContable.ClaseCuenta.RESULTADO_POSITIVO
            if 'negativo' in c_low or 'egreso' in c_low or 'gasto' in c_low or 'costo' in c_low:
                return CuentaContable.ClaseCuenta.RESULTADO_NEGATIVO
            
            # Fallback por prefijo contable
            if code.startswith('1.'):
                return CuentaContable.ClaseCuenta.ACTIVO
            if code.startswith('2.'):
                return CuentaContable.ClaseCuenta.PASIVO
            if code.startswith('3.'):
                return CuentaContable.ClaseCuenta.PATRIMONIO_NETO
            if code.startswith('4.1.'):
                return CuentaContable.ClaseCuenta.RESULTADO_POSITIVO
            if code.startswith('4.2.'):
                return CuentaContable.ClaseCuenta.RESULTADO_NEGATIVO
            return CuentaContable.ClaseCuenta.OTRO

        def normalize_saldo(saldo_str):
            s_low = saldo_str.lower()
            if 'deudor' in s_low:
                return CuentaContable.SaldoHabitual.DEUDOR
            if 'acreedor' in s_low:
                return CuentaContable.SaldoHabitual.ACREEDOR
            return ''

        with transaction.atomic():
            if options['limpiar']:
                self.stdout.write(self.style.WARNING("[INFO] Limpiando cuentas contables previas..."))
                CuentaContable.objects.all().delete()

            creadas = 0
            actualizadas = 0
            cuentas_map = {}

            # Paso 1: Crear o actualizar cuentas sin asignar padre todavía
            for row in raw_rows:
                codigo = row['codigo']
                nombre = row['nombre']
                saldo = normalize_saldo(row['saldo'])
                clase = normalize_class(row['clase'], codigo)
                nivel = determine_level(codigo)
                es_imputable = bool(saldo != '')
                nota = row['nota']

                cuenta, created = CuentaContable.objects.update_or_create(
                    codigo=codigo,
                    defaults={
                        'nombre': nombre,
                        'saldo_habitual': saldo,
                        'clase': clase,
                        'nivel': nivel,
                        'es_imputable': es_imputable,
                        'nota': nota,
                        'activa': True
                    }
                )
                cuentas_map[codigo] = cuenta
                if created:
                    creadas += 1
                else:
                    actualizadas += 1

            # Paso 2: Asignar cuentas padre
            padres_asignados = 0
            for row in raw_rows:
                codigo = row['codigo']
                parent_code = find_parent_code(codigo)
                if parent_code and parent_code in cuentas_map:
                    cuenta = cuentas_map[codigo]
                    cuenta.padre = cuentas_map[parent_code]
                    cuenta.save(update_fields=['padre'])
                    padres_asignados += 1

            self.stdout.write(
                self.style.SUCCESS(
                    f"[OK] Proceso completado: {creadas} creadas, {actualizadas} actualizadas. "
                    f"{padres_asignados} relaciones jerarquicas vinculadas."
                )
            )

            # Paso 3: Enlace y Auto-mapeo con modelos existentes
            if not options['sin_mapeo']:
                self.mapear_modelos_existentes(cuentas_map)

    def mapear_modelos_existentes(self, cuentas_map):
        self.stdout.write("\n[VINCULACION] Iniciando auto-mapeo con Cajas, Insumos y Centros de Costo...")

        # 1. Cuentas de Tesorería (Cajas y Bancos)
        for cuenta_fin in Cuenta.objects.all():
            nombre_l = cuenta_fin.nombre.lower()
            cuenta_contable = None

            if 'dólar' in nombre_l or 'usd' in nombre_l or cuenta_fin.moneda == 'USD':
                if cuenta_fin.tipo == Cuenta.TipoCuenta.CUENTA_BANCARIA:
                    cuenta_contable = cuentas_map.get('1.1.1.02.000002')  # Banco cta. cte. en moneda extranjera
                else:
                    cuenta_contable = cuentas_map.get('1.1.1.01.000003')  # Caja en Moneda Extranjera
            elif cuenta_fin.tipo == Cuenta.TipoCuenta.CUENTA_BANCARIA or 'banco' in nombre_l:
                cuenta_contable = cuentas_map.get('1.1.1.02.000001')  # Banco cta. cte. en pesos
            elif 'chica' in nombre_l:
                cuenta_contable = cuentas_map.get('1.1.1.01.000002')  # Caja Chica
            else:
                cuenta_contable = cuentas_map.get('1.1.1.01.000001')  # Caja Central

            if cuenta_contable:
                cuenta_fin.cuenta_contable = cuenta_contable
                cuenta_fin.save(update_fields=['cuenta_contable'])
                self.stdout.write(f"   [Caja/Banco] '{cuenta_fin.nombre}' -> [{cuenta_contable.codigo}] {cuenta_contable.nombre}")

        # 2. Categorías de Insumos (Activo y Consumo/Gasto)
        for cat in CategoriaInsumo.objects.all():
            cat_l = cat.nombre.lower()
            cta_activo = None
            cta_gasto = None

            if 'fertiliz' in cat_l:
                cta_activo = cuentas_map.get('1.1.5.01.000001')  # Fertilizantes
                cta_gasto = cuentas_map.get('4.2.1.02.000000')   # Fertilizantes, agroquímicos y riego
            elif 'herbic' in cat_l:
                cta_activo = cuentas_map.get('1.1.5.01.000002')  # Herbicidas
                cta_gasto = cuentas_map.get('4.2.1.02.000000')
            elif any(w in cat_l for w in ['insectic', 'fungic', 'agroquím', 'fito']):
                cta_activo = cuentas_map.get('1.1.5.01.000003')  # Insecticidas / Fungicidas
                cta_gasto = cuentas_map.get('4.2.1.02.000000')
            elif 'riego' in cat_l:
                cta_activo = cuentas_map.get('1.1.5.01.000004')  # Riego - insumos
                cta_gasto = cuentas_map.get('4.2.1.02.000000')
            elif any(w in cat_l for w in ['envase', 'botella', 'tapa', 'etiqueta', 'embalaje']):
                cta_activo = cuentas_map.get('1.1.5.05.000001') or cuentas_map.get('1.1.5.05.000000')  # Botellas / Envases
                cta_gasto = cuentas_map.get('4.2.3.02.000000')   # Envases, tapas y etiquetas consumidos
            else:
                # Default agrícola
                cta_activo = cuentas_map.get('1.1.5.01.000000')  # Insumos y materiales producción agrícola
                cta_gasto = cuentas_map.get('4.2.1.02.000000')

            changed = False
            if cta_activo:
                cat.cuenta_contable_activo = cta_activo
                changed = True
            if cta_gasto:
                cat.cuenta_contable_gasto = cta_gasto
                changed = True
            if changed:
                cat.save()
                self.stdout.write(f"   [Categoria Insumo] '{cat.nombre}' -> Activo: [{cta_activo.codigo if cta_activo else '-'}] | Gasto: [{cta_gasto.codigo if cta_gasto else '-'}]")

        # 3. Centros de Costo
        for cc in CentroDeCosto.objects.all():
            cta_defecto = None
            if cc.tipo == 'PRODUCTIVO_CAMPO':
                cta_defecto = cuentas_map.get('4.2.1.00.000000')  # Costo de la producción agrícola
            elif cc.tipo == 'FABRICA_ALMAZARA':
                cta_defecto = cuentas_map.get('4.2.2.00.000000')  # Costo producción industrial - elaboración de aceite
            elif cc.tipo == 'MAQUINARIA_TALLER':
                cta_defecto = cuentas_map.get('4.2.1.07.000000') or cuentas_map.get('4.2.2.03.000000')
            elif cc.tipo == 'ESTRUCTURA_ADMIN':
                cta_defecto = cuentas_map.get('4.2.5.00.000000')  # Gastos de administración
            elif cc.tipo == 'COMERCIAL_EXPORT':
                cta_defecto = cuentas_map.get('4.2.6.00.000000')  # Gastos de comercialización

            if cta_defecto:
                cc.cuenta_contable_defecto = cta_defecto
                cc.save(update_fields=['cuenta_contable_defecto'])
                self.stdout.write(f"   [Centro de Costo] '{cc.nombre}' ({cc.get_tipo_display()}) -> [{cta_defecto.codigo}] {cta_defecto.nombre}")

        self.stdout.write(self.style.SUCCESS("[OK] Mapeo automatico de entidades completado con exito."))
